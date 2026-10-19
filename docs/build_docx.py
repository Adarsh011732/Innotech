import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))

FILES = [
    ("1_PRD.md", "1_PRD.docx"),
    ("2_TECHNICAL_ARCHITECTURE.md", "2_TECHNICAL_ARCHITECTURE.docx"),
    ("3_SECURITY_AND_ACCESS.md", "3_SECURITY_AND_ACCESS.docx"),
    ("4_FRONTEND_SPEC.md", "4_FRONTEND_SPEC.docx"),
    ("5_FEATURE_TICKET_LIST.md", "5_FEATURE_TICKET_LIST.docx"),
]

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(14, 116, 144) # Deep Teal / Cyan
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(31, 41, 55)   # Dark Charcoal
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(55, 65, 81)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    pPr = p._element.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    pPr.append(shd)

def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF") # Soft Blue
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 58, 138)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def parse_inline_formatting(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 41, 55)
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.italic = True
            run.font.color.rgb = RGBColor(55, 65, 81)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 30, 30)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(55, 65, 81)

def render_markdown_table(doc, lines):
    rows_data = []
    for line in lines:
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if any(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, color="D1D5DB", sz="4", val="single")
    
    for r_idx, row in enumerate(rows_data):
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            
            if r_idx == 0:
                set_cell_background(cell, "1E3A8A") # Navy Blue header
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                run = p.add_run(val)
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                parse_inline_formatting(p, val)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.startswith('```'):
            if in_code_block:
                add_code_block(doc, '\n'.join(code_lines))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Check table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line.strip())
            in_table = True
            i += 1
            continue
        elif in_table:
            render_markdown_table(doc, table_lines)
            table_lines = []
            in_table = False
            
        # Callouts / blockquotes
        if line.startswith('> '):
            add_callout(doc, line[2:].strip())
            i += 1
            continue
            
        # Horizontal rule
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._element.get_or_add_pPr()
            pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            pPr.append(pbdr)
            i += 1
            continue
            
        # Headings
        if line.startswith('# '):
            add_styled_heading(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith('## '):
            add_styled_heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith('### '):
            add_styled_heading(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith('#### '):
            add_styled_heading(doc, line[5:].strip(), 4)
            i += 1
            continue
            
        # Bullet list
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            indent_level = (len(line) - len(line.lstrip())) // 2
            bullet_text = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parse_inline_formatting(p, bullet_text)
            i += 1
            continue
            
        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)', line.strip())
        if num_match:
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parse_inline_formatting(p, item_text)
            i += 1
            continue
            
        # Normal paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, line.strip())
            
        i += 1
        
    if in_table and table_lines:
        render_markdown_table(doc, table_lines)
        
    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    for md_name, docx_name in FILES:
        md_file = os.path.join(DOCS_DIR, md_name)
        docx_file = os.path.join(DOCS_DIR, docx_name)
        if os.path.exists(md_file):
            convert_md_to_docx(md_file, docx_file)
        else:
            print(f"File not found: {md_file}")
